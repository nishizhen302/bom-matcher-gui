# -*- coding: utf-8 -*-
"""生产版输出（finalize）测试：模拟"导出→人工审核→生成生产版"全流程。"""
import os
import sys

import openpyxl
import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from bom_matcher import report
from bom_matcher.finalize import finalize, FINAL_HEADER


@pytest.fixture
def reviewed_file(tmp_path):
    """先用 report.write_excel 生成真实格式的结果文件，再模拟人工审核编辑。"""
    rows = [
        ['C1', '0.1uF', 'C0603', '1', '100-000123', '贴片电容', '0.1uF 50V X7R 0603', '', ''],
        ['C2', '0.1uF', 'C0603', '2', '100-000123', '贴片电容', '0.1uF 50V X7R 0603', '', ''],
        ['R1', '10K', 'R0603', '1', '100-000200', '贴片电阻', '10K 1% 0603', '', ''],
        ['H1', '安装孔', 'HOLE', '1', report.MECH_PN, '安装孔/测试点', '', '', '非采购物料，不参与匹配'],
        ['D5', 'NC', 'C0603', '1', report.UNMATCHED, report.UNMATCHED, report.UNMATCHED, '', '编码表无对应料'],
        ['C9', '0.1uF', 'C0805', '1', report.UNMATCHED, report.UNMATCHED, report.UNMATCHED, '', '编码表无对应料'],
        ['X1', '22pF', 'C0402', '1', report.UNMATCHED, report.UNMATCHED, report.UNMATCHED, '', '编码表无对应料'],
        ['AUTO', '锡膏', '', '3.2g', '100-000900', '锡膏', '有铅锡膏', '', '联动规则自动追加'],
    ]
    report.sort_rows(rows)
    result = str(tmp_path / '结果.xlsx')
    report.write_excel(rows, result)

    # 模拟人工审核：C9 手填料号（黄色底色还在）；追加一行 PCB 料号（无位号）
    wb = openpyxl.load_workbook(result)
    ws = wb.active
    for row in ws.iter_rows(min_row=2):
        if str(row[0].value) == 'C9':
            row[4].value, row[5].value, row[6].value = '100-000123', '贴片电容', '0.1uF 50V X7R 0603'
    pcb = ws.max_row + 1
    for col, val in ((4, '1'), (5, '100-000PCB'), (6, 'PCB'), (7, '4层 FR4 1.6mm')):
        ws.cell(row=pcb, column=col).value = val
    wb.save(result)
    return result


def _load(path):
    ws = openpyxl.load_workbook(path).active
    data = [[c.value for c in r] for r in ws.iter_rows(min_row=2)]
    return ws, data


def test_finalize(reviewed_file):
    out = finalize(reviewed_file, log=lambda m: None)
    assert out.endswith('_生产版.xlsx') and os.path.exists(out)
    ws, data = _load(out)

    assert [c.value for c in ws[1]] == FINAL_HEADER
    assert ws.freeze_panes == 'A2'
    assert ws.max_column == len(FINAL_HEADER)

    # 手动追加的 PCB 行排第一，序号从 1 连续编号
    assert data[0][1] == '100-000PCB' and data[0][0] == 1
    assert [r[0] for r in data] == list(range(1, len(data) + 1))

    by_pn = {r[1]: r for r in data}
    # C9 手填料号后并入同料号汇总：数量 1+2+1，位号合并
    agg = by_pn['100-000123']
    assert agg[4] == '4'
    assert sorted(agg[5].split(',')) == ['C1', 'C2', 'C9']
    assert by_pn['100-000200'][4] == '1'
    # 锡膏行不聚合，原样保留
    assert by_pn['100-000900'][4] == '3.2g'

    # NC 行、非采购项行、旧汇总行都不出现
    assert report.MECH_PN not in by_pn
    refs = ','.join(str(r[5]) for r in data)
    assert 'D5' not in refs and 'H1' not in refs

    # 漏填的 X1 保留在表尾且为黄色
    assert data[-1][1] == report.UNMATCHED and data[-1][5] == 'X1'
    last = list(ws.iter_rows(min_row=ws.max_row))[0]
    assert str(last[0].fill.start_color.rgb).endswith('FFFF00')
    # 其余行无底色
    first = list(ws.iter_rows(min_row=2, max_row=2))[0]
    assert first[0].fill.fill_type != 'solid'


def test_finalize_warns_leftover(reviewed_file, capsys):
    logs = []
    finalize(reviewed_file, log=logs.append)
    assert any('X1' in m and '未填料号' in m for m in logs)


def test_finalize_rejects_foreign_file(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.append(['随便', '什么', '表'])
    path = str(tmp_path / 'foreign.xlsx')
    wb.save(path)
    with pytest.raises(ValueError):
        finalize(path, log=lambda m: None)


class _E:
    def __init__(self, name, spec):
        self.name, self.spec = name, spec


def test_pick_candidate_by_dropdown(tmp_path):
    """2026-10-18：带候选的行料号格可下拉换料，生产版按选中料号补齐名称/规格（不用手抄）。"""
    catalog = {'Y0604-00010': _E('贴片电感', 'SMPA0503-4R7MT,4.7uH'),
               'Y0604-00033': _E('贴片电感', 'SLO0520H4R7MTT,4.7uH'),
               'Y0612-00114': _E('IC', 'LT87121,VBO转DP')}
    rows = [
        ['L1', '4.7uH', 'IND', '1', 'Y0604-00010', '贴片电感', 'SMPA0503-4R7MT,4.7uH',
         'Y0604-00033: SLO0520H4R7MTT,4.7uH', '感值4.7uH [需确认]'],
        ['L2', '4.7uH', 'IND', '1', 'Y0604-00010', '贴片电感', 'SMPA0503-4R7MT,4.7uH',
         'Y0604-00033: SLO0520H4R7MTT,4.7uH', '感值4.7uH [需确认]'],
        ['U3', 'LT87121', 'QFN64', '1', report.UNMATCHED, report.UNMATCHED, report.UNMATCHED,
         'Y0612-00114: LT87121,VBO转DP; Y9999-99999: 编码表里没有', '编码表无对应料'],
        ['R1', '10K', 'R0402', '1', 'Y0601-00050', '贴片电阻', '10K', '', ''],
    ]
    report.sort_rows(rows)
    result = str(tmp_path / '结果.xlsx')
    report.write_excel(rows, result, catalog=catalog)

    wb = openpyxl.load_workbook(result)
    ws = wb.worksheets[0]
    assert wb[report.CATALOG_SHEET].sheet_state == 'hidden'
    dv_of = {}
    for dv in ws.data_validations.dataValidation:
        for rng in dv.sqref.ranges:
            dv_of[rng.coord] = dv.formula1
    row_of = {str(r[0].value): r[0].row for r in ws.iter_rows(min_row=2)}
    assert dv_of[f'E{row_of["L1"]}'] == '"Y0604-00010,Y0604-00033"'   # 当前料号排第一
    assert dv_of[f'E{row_of["U3"]}'] == '"Y0612-00114"'               # 编码表里没有的候选不进下拉
    assert f'E{row_of["R1"]}' not in dv_of                             # 没候选的行不加下拉

    # 模拟在 Excel 里下拉：L1 换候选、U3 选候选，名称/规格都没动
    ws.cell(row=row_of['L1'], column=5).value = 'Y0604-00033'
    ws.cell(row=row_of['U3'], column=5).value = 'Y0612-00114'
    wb.save(result)

    logs = []
    out = finalize(result, log=logs.append)
    got = {r[1]: r for r in openpyxl.load_workbook(out).active.iter_rows(min_row=2, values_only=True)}
    assert got['Y0604-00033'][2:6] == ('贴片电感', 'SLO0520H4R7MTT,4.7uH', '1', 'L1')
    assert got['Y0604-00010'][3:6] == ('SMPA0503-4R7MT,4.7uH', '1', 'L2')
    assert got['Y0612-00114'][2:4] == ('IC', 'LT87121,VBO转DP')
    assert any(m.startswith('2 行按料号') for m in logs)


def test_handfilled_pn_looked_up_in_code_table(tmp_path):
    """2026-10-18：没匹配上的行只手填料号（名称还是「未匹配」），生产版按编码表补名称/规格。"""
    rows = [
        ['U3', 'LT87121', 'QFN64', '1', report.UNMATCHED, report.UNMATCHED, report.UNMATCHED, '', '编码表无对应料'],
        ['U4', 'XX', 'SOT23', '1', report.UNMATCHED, report.UNMATCHED, report.UNMATCHED, '', '编码表无对应料'],
        ['U5', 'YY', 'SOT23', '1', report.UNMATCHED, report.UNMATCHED, report.UNMATCHED, '', '编码表无对应料'],
    ]
    result = str(tmp_path / '结果.xlsx')
    report.write_excel(rows, result)
    wb = openpyxl.load_workbook(result)
    ws = wb.worksheets[0]
    typed = {'U3': ' y0612-00114 ', 'U4': 'Y0601-00141', 'U5': 'Y0000-00000'}  # 小写带空格/停用料/编码表没有
    for r in ws.iter_rows(min_row=2):
        r[4].value = typed[str(r[0].value)]
    wb.save(result)

    logs = []
    out = finalize(result, log=logs.append,
                   code_lookup={'Y0612-00114': ('集成电路', 'LT87121，VBO转DP')},
                   stopped={'Y0601-00141'})
    got = {r[5]: r for r in openpyxl.load_workbook(out).active.iter_rows(min_row=2, values_only=True)}
    assert got['U3'][1:4] == ('Y0612-00114', '集成电路', 'LT87121，VBO转DP')
    assert any('停用料' in m and 'U4' in m for m in logs)
    assert any('查不到' in m and 'U5' in m and 'U3' not in m for m in logs)
