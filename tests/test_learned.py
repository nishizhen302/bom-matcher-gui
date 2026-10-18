# -*- coding: utf-8 -*-
"""人工选择记忆：改过一次，下次同样写法直接配上。"""
import os
import sys

import openpyxl
import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from bom_matcher import learned as LR, report
from bom_matcher.engine import Engine
from bom_matcher.finalize import collect_manual_picks
from bom_matcher.loader import _build_entry

CODE = [
    ('Y0610-00001', 'MOS管', 'FDS4435BZ,单,P,30V,8.8A,150℃,SOP-8,ON'),
    ('Y0606-00006', '磁珠', 'BLM18EG101TN1D,100Ω@100MHz，2A，0603，Murata'),
    ('Y0612-00105', '集成电路', 'SGM61431,36V,3A,SOIC-8N,SGMICRO'),
]


def _eng(store):
    return Engine([_build_entry(*c) for c in CODE], learned=store)


def test_key_ignores_case_separators_and_nc():
    assert LR.key_of('4435', 'SOIC8') == LR.key_of('4435', 'soic-8')
    assert LR.key_of('NC/4435', 'SOIC8') == LR.key_of('4435', 'SOIC8')
    assert LR.key_of('4435', 'SOIC8') != LR.key_of('4435', 'SOT23')   # 封装不同算不同写法


def test_learned_pick_used_next_time_and_beats_rules():
    store = {}
    assert LR.remember(store, '4435', 'SOIC127P600X175-8N', 'Y0610-00001', board='2795')
    eng = _eng(store)
    m = eng.match(eng.build_item('U2', '4435', 'SOIC127P600X175-8N', '1'))
    assert m.pn == 'Y0610-00001' and '沿用人工选择' in m.reason and eng.learned_used == 1
    # 比内置规则优先：磁珠规则本来给 Y0606-00006，用户指定过别的就用用户的
    LR.remember(store, '220R/100MHz', 'L0603', 'Y0612-00105')
    eng = _eng(store)
    assert eng.match(eng.build_item('FB8', '220R/100MHz', 'L0603', '1')).pn == 'Y0612-00105'
    # 同样的 Part 但封装不同：不沿用
    eng = _eng(store)
    assert eng.match(eng.build_item('U9', '4435', 'SOT23', '1')) is None


def test_learned_pn_gone_from_code_table_is_skipped():
    store = {}
    LR.remember(store, 'RClamp0564P', 'RClamp0564P', 'Y0615-99999')   # 后来编码表里没有了
    eng = _eng(store)
    assert eng.match(eng.build_item('UX5', 'RClamp0564P', 'RClamp0564P', '1')) is None
    assert eng.learned_skipped and eng.learned_skipped[0][2] == 'Y0615-99999'


def test_remember_overwrites_with_latest_and_survives_reload(tmp_path):
    store, p = {}, str(tmp_path / 'learned.json')
    LR.remember(store, '4435', 'SOIC8', 'Y0610-00001')
    assert not LR.remember(store, '4435', 'SOIC8', 'Y0610-00001')    # 一样的不算变化
    assert LR.remember(store, '4435', 'SOIC8', 'Y0610-00002')        # 以最新为准
    LR.save(p, store)
    assert LR.load(p)[LR.key_of('4435', 'SOIC8')]['pn'] == 'Y0610-00002'
    assert LR.load(str(tmp_path / 'nope.json')) == {}


def _result_with_edits(tmp_path):
    rows = [
        ['U2', '4435', 'SOIC8', '1', report.UNMATCHED, report.UNMATCHED, report.UNMATCHED, '', '编码表无对应料'],
        ['U4,U8', 'SGM61430', 'SOIC8', '2', 'Y0612-00105', '集成电路', 'SGM61431', '', 'IC近似 [需确认]'],
        ['FB1', '100R', 'L0603', '1', 'Y0606-00006', '磁珠', 'BLM18EG101TN1D', '', ''],
        ['C9', 'NC/100nF', 'C0402', '1', report.UNMATCHED, report.UNMATCHED, report.UNMATCHED, '', ''],
    ]
    report.sort_rows(rows)
    f = str(tmp_path / '结果.xlsx')
    report.write_excel(rows, f)
    wb = openpyxl.load_workbook(f)
    ws = wb.worksheets[0]
    for r in ws.iter_rows(min_row=2):
        ref = str(r[0].value)
        if ref == 'U2':
            r[4].value = 'y0610-00001'       # 手填料号（小写）
        elif ref == 'C9':
            r[4].value = 'Y0602-00068'       # NC 行改了也不记
    wb.save(f)
    return f


def test_collect_manual_picks_from_excel_edits(tmp_path):
    picks = collect_manual_picks(_result_with_edits(tmp_path))
    assert [(p[0], p[1], p[2]) for p in picks] == [('4435', 'SOIC8', 'y0610-00001')]   # 没改的行、NC 行都不算


def test_old_result_without_orig_sheet_learns_nothing(tmp_path):
    f = _result_with_edits(tmp_path)
    wb = openpyxl.load_workbook(f)
    del wb[report.ORIG_SHEET]
    wb.save(f)
    assert collect_manual_picks(f) == []


def test_gui_learns_on_finalize_and_review(tmp_path, monkeypatch):
    import app
    import excel_bridge
    p = str(tmp_path / 'learned.json')
    monkeypatch.setattr(app, 'LEARNED_FILE', p)
    monkeypatch.setattr(excel_bridge, 'status', lambda f: excel_bridge.CLOSED)
    a = app.Api()
    a.final_file = _result_with_edits(tmp_path)
    monkeypatch.setattr(a, '_code_lookup', lambda f: ({'Y0610-00001': ('MOS管', 'FDS4435BZ')}, set()))
    a.code_file = __file__   # 只要文件存在；查表已替身
    r = a.make_final('check')
    assert r.get('ok') and '记住了 1 条' in r['info']
    assert a.learned_count() == 1
    rec = a.list_learned()[0]
    assert (rec['part'], rec['pn'], rec['name']) == ('4435', 'Y0610-00001', 'MOS管')
    assert a.forget_learned(rec['key'])['count'] == 0
    # 复核里「采用」也记
    a.session = {'review_rows': [['UX5', 'RClamp0564P', 'RClamp0564P', '1', '未匹配', '', '']],
                 'export': lambda res: None, 'results': [], 'dtype_map': {}}
    a.bom_file = 'C:/pcb/2795/RTD2795.DSN'
    a.apply_override(0, 'Y0615-00001', '静电保护器', 'RCLAMP0524P')
    assert LR.load(p)[LR.key_of('RClamp0564P', 'RClamp0564P')]['board'] == 'RTD2795'
