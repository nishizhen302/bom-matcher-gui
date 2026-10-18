# -*- coding: utf-8 -*-
"""输出层：排序、同料号聚合、汇总行、着色、列宽行高。格式与旧版一致，新增第9列"匹配说明"。"""
import math
import re
import openpyxl

from .parsers import NC_RE

UNMATCHED = '未匹配'
MECH_PN = '非采购项'

HEADER = ['位号', 'Part Value', 'PCB Footprint', 'Quantity', '料号', '名称', '规格', '候选料号', '匹配说明']
N_COLS = len(HEADER)

# 隐藏页：带候选的行里出现的每个料号的完整名称/规格。
# 「料号」格下拉换成候选后，生产版据此补齐名称/规格（Excel 里不用手抄）
CATALOG_SHEET = '候选料号库'
ORIG_SHEET = '引擎原结果'


def candidate_pns(r):
    """候选列 'PN: spec; PN: spec' → [PN...]"""
    if len(r) <= 7 or not r[7]:
        return []
    return [c.split(':', 1)[0].strip() for c in str(r[7]).split('; ') if c.strip()]


def _is_nc(part):
    return bool(NC_RE.search(str(part).upper()))


def _is_review(r):
    return len(r) > 8 and '[需确认]' in str(r[8])


def sort_rows(rows):
    """稳定排序：正常(按料号) → [需确认] → 非采购项 → 未匹配 → NC。
    带颜色的行集中放在表尾一个块里，便于逐块人工处理。"""
    def tier(r):
        if _is_nc(r[1]):
            return 4
        if r[4] == UNMATCHED:
            return 3
        if r[4] == MECH_PN:
            return 2
        if _is_review(r):
            return 1
        return 0

    rows.sort(key=lambda r: (tier(r), r[4] == UNMATCHED, r[4]))
    return rows


def _parse_qty(value):
    m = re.search(r'\d+(?:\.\d+)?', str(value or ''))
    return float(m.group(0)) if m else 0.0


def add_companions(rows, companions, by_pn, log):
    for trigger_pn, (add_pn, per_qty, label) in companions.items():
        total = sum(_parse_qty(r[3]) for r in rows if r[4] == trigger_pn)
        if total <= 0 or any(r[4] == add_pn for r in rows):
            continue
        e = by_pn.get(add_pn)
        if not e:
            log(f"Warning: 匹配到 {trigger_pn}，但编码表未找到可用的 {add_pn}")
            continue
        qty_text = f"{total * per_qty:.3f}".rstrip('0').rstrip('.') + 'g'
        rows.append(['AUTO', label, '', qty_text, e.pn, e.name, e.spec, '', '联动规则自动追加'])
        log(f"自动添加 {add_pn}: {total:g} * {per_qty}g = {qty_text}")
    return rows


def _add_pick_lists(wb, ws, data_rows, catalog):
    """带候选的明细行：料号格加下拉（当前料号+候选），候选的名称/规格写进隐藏页。
    不拦手输——编码表里有、但没进候选的料仍可直接敲。"""
    from openpyxl.worksheet.datavalidation import DataValidation
    lists = {}      # 下拉内容 -> [行号]
    known = {}      # 料号 -> (名称, 规格)
    for row_num, r in data_rows:
        cands = [p for p in candidate_pns(r) if p in catalog]  # 规格里带 '; ' 时切出的碎片不算
        if not cands:
            continue
        opts = ([r[4]] if r[4] not in (UNMATCHED, MECH_PN, '') else []) + \
               [p for p in cands if p != r[4]]
        # Excel 内联下拉上限 255 字符，超出截掉末尾候选
        while opts and len(','.join(opts)) > 250:
            opts.pop()
        if not opts:
            continue
        lists.setdefault(','.join(opts), []).append(row_num)
        for pn in opts:
            e = catalog.get(pn)
            if e is not None:
                known[pn] = (e.name, e.spec)
    if not lists:
        return
    for opts, row_nums in lists.items():
        dv = DataValidation(type='list', formula1=f'"{opts}"', allow_blank=True,
                            showErrorMessage=False, showInputMessage=True,
                            promptTitle='选料号', prompt='下拉可换成候选料号；名称/规格在生成生产版时自动补齐')
        ws.add_data_validation(dv)
        for n in row_nums:
            dv.add(f'E{n}')
    cs = wb.create_sheet(CATALOG_SHEET)
    cs.append(['料号', '名称', '规格'])
    for pn in sorted(known):
        cs.append([pn, *known[pn]])
    cs.sheet_state = 'hidden'


def read_catalog(wb):
    """结果文件里的隐藏候选库 → {料号: (名称, 规格)}；旧版结果文件没有就返回 {}。"""
    if CATALOG_SHEET not in wb.sheetnames:
        return {}
    out = {}
    for pn, name, spec in wb[CATALOG_SHEET].iter_rows(min_row=2, max_col=3, values_only=True):
        if pn:
            out[str(pn).strip()] = (str(name or ''), str(spec or ''))
    return out


def write_excel(rows, output_file, catalog=None):
    """catalog: 料号 -> 编码表条目（有 .name/.spec），给出时为带候选的行生成下拉选料。"""
    aggregatable = lambda r: (not _is_nc(r[1])) and r[4] not in (UNMATCHED, MECH_PN)
    pn_counts = {}
    for r in rows:
        if aggregatable(r):
            pn_counts[r[4]] = pn_counts.get(r[4], 0) + 1
    repeated = {pn for pn, c in pn_counts.items() if c > 1}

    write_rows = []  # (type, data, is_repeated)
    group = []

    def flush(g):
        if len(g) > 1 and aggregatable(g[0]):
            total = sum(int(str(x[3])) for x in g if str(x[3]).isdigit())
            refs = ','.join(str(x[0]) for x in g)
            write_rows.append(('summary',
                               [refs, '', '', str(total), g[0][4], g[0][5], g[0][6], '', ''], True))

    for r in rows:
        if not group or group[-1][4] != r[4]:
            flush(group)
            group = [r]
        else:
            group.append(r)
        write_rows.append(('data', r, r[4] in repeated))
    flush(group)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(HEADER)
    thin = openpyxl.styles.Side(style='thin')
    border = openpyxl.styles.Border(left=thin, right=thin, top=thin, bottom=thin)
    fills = {
        'yellow': openpyxl.styles.PatternFill('solid', start_color='FFFF00'),
        'blue': openpyxl.styles.PatternFill('solid', start_color='0099CC'),
        'gray': openpyxl.styles.PatternFill('solid', start_color='D9D9D9'),
        'green': openpyxl.styles.PatternFill('solid', start_color='C6EFCE'),
        'mech': openpyxl.styles.PatternFill('solid', start_color='FFF2CC'),
        'review': openpyxl.styles.PatternFill('solid', start_color='F8CBAD'),
    }
    for col in range(1, N_COLS + 1):
        ws.cell(row=1, column=col).border = border

    data_rows = []
    for row_type, data, is_rep in write_rows:
        row_num = ws.max_row + 1
        ws.append(data[:N_COLS] + [''] * (N_COLS - len(data)))
        if row_type == 'data':
            data_rows.append((row_num, data))
        fill = None
        if row_type == 'summary':
            fill = fills['green']
            for col in range(1, N_COLS + 1):
                ws.cell(row=row_num, column=col).font = openpyxl.styles.Font(bold=True)
        elif _is_nc(data[1]):
            fill = fills['blue']
        elif data[4] == UNMATCHED:
            fill = fills['yellow']
        elif len(data) > 8 and '[需确认]' in str(data[8]):
            fill = fills['review']
        elif data[4] == MECH_PN:
            fill = fills['mech']
        elif is_rep:
            fill = fills['gray']
        for col in range(1, N_COLS + 1):
            cell = ws.cell(row=row_num, column=col)
            cell.border = border
            if fill:
                cell.fill = fill

    # 对齐
    center = openpyxl.styles.Alignment(horizontal='center', vertical='center', wrap_text=True)
    left = openpyxl.styles.Alignment(horizontal='left', vertical='center', wrap_text=True)
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
        for cell in row:
            cell.alignment = center if (cell.row == 1 or cell.column in (4, 5, 6)) else left

    widths = {1: 44, 2: 16.5, 3: 16.5, 4: 10, 5: 16, 6: 14, 7: 44, 8: 11, 9: 22}
    for col, w in widths.items():
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = w

    ws.row_dimensions[1].height = 22
    col_eff = {1: 36, 2: 13, 3: 13, 7: 36, 8: 9, 9: 18}
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        max_lines = 1
        for cell in row:
            if cell.column in col_eff and cell.value:
                w = sum(2 if ord(c) > 127 else 1 for c in str(cell.value))
                max_lines = max(max_lines, math.ceil(w / col_eff[cell.column]))
        ws.row_dimensions[row[0].row].height = max(30, max_lines * 15)

    if catalog:
        _add_pick_lists(wb, ws, data_rows, catalog)
    # 隐藏页存匹配时的原料号：生成生产版时比对出用户在 Excel 里改过哪些，记进人工选择
    os_ = wb.create_sheet(ORIG_SHEET)
    os_.append(['位号', 'Part', 'PCB Footprint', '料号'])
    for _n, d in data_rows:
        os_.append([d[0], d[1], d[2], d[4]])
    os_.sheet_state = 'hidden'
    wb.save(output_file)
