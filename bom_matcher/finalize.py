# -*- coding: utf-8 -*-
"""生产版输出：读回人工审核后的匹配结果 Excel，删 NC/非采购项/旧汇总行，
按料号重新聚合排序，输出 6 列（序号/料号/名称/规格/数量/位号）统一格式的 <原名>_生产版.xlsx。

行分类规则（针对本工具导出的 9 列结果文件）：
  - 旧汇总行：绿色填充或加粗 —— 丢弃，明细重新聚合
  - NC 行 / 料号="非采购项" 行 —— 丢弃
  - 位号为空的行：人工追加物料（如 PCB 板料号），保留并按原顺序排在最前
  - 料号仍为"未匹配"的行：漏填，保留黄色置于表尾并在日志中警告
  - 数量非纯数字的行（如锡膏 x.xg）：不参与聚合，原样保留在表尾
"""
import math
import os
import re

import openpyxl

from .parsers import NC_RE
from .report import HEADER, UNMATCHED, MECH_PN, ORIG_SHEET, read_catalog

FINAL_HEADER = ['序号', '料号', '名称', '规格', '数量', '位号']
SUMMARY_GREEN = 'C6EFCE'


def _is_summary(cells):
    for c in cells:
        if c.font and c.font.bold:
            return True
        fill = c.fill
        if fill and fill.fill_type == 'solid' and str(fill.start_color.rgb or '').endswith(SUMMARY_GREEN):
            return True
    return False


def read_result(result_file, code_lookup=None):
    """读回结果 Excel，分类为 (manual, details, specials, leftovers, n_filled)。
    detail 行格式统一为 [位号, 料号, 名称, 规格, 数量]。
    名称/规格以料号为准补齐：先查编码表 code_lookup（手填料号也能对上），
    再查结果文件里的隐藏候选库（下拉换的料）；n_filled 为实际改动行数。"""
    wb = openpyxl.load_workbook(result_file)
    ws = wb.worksheets[0]
    catalog = read_catalog(wb)
    catalog.update(code_lookup or {})
    n_filled = 0
    header = [str(c.value or '').strip() for c in ws[1][:len(HEADER)]]
    if header[:5] != HEADER[:5]:
        raise ValueError(f"表头不是本工具的匹配结果格式（期望 {HEADER[:5]}...，实际 {header[:5]}）")

    manual, details, specials, leftovers = [], [], [], []
    for row in ws.iter_rows(min_row=2):
        cells = row[:len(HEADER)]
        v = [str(c.value or '').strip() for c in cells]
        v += [''] * (len(HEADER) - len(v))
        ref, part, qty, pn, name, spec = v[0], v[1], v[3], v[4], v[5], v[6]
        if not any(v):
            continue
        if _is_summary(cells):
            continue
        if NC_RE.search(part.upper()):
            continue
        if pn == MECH_PN:
            continue
        if pn.upper() in catalog:      # 手敲小写也认
            pn = pn.upper()
        if pn in catalog and (name, spec) != catalog[pn]:
            name, spec = catalog[pn]
            n_filled += 1
        item = [ref, pn, name, spec, qty]
        if pn == UNMATCHED:
            leftovers.append(item)
        elif not ref:
            manual.append(item)
        elif not pn:
            leftovers.append(item)
        elif re.fullmatch(r'\d+', qty):
            details.append(item)
        else:
            specials.append(item)
    return manual, details, specials, leftovers, n_filled


def collect_manual_picks(result_file):
    """用户在 Excel 里改过料号的行（对比隐藏页「引擎原结果」）→ [(Part, 封装, 料号, 位号)]。
    旧版结果文件没有该隐藏页：返回 []（不猜）。"""
    wb = openpyxl.load_workbook(result_file)
    if ORIG_SHEET not in wb.sheetnames:
        return []
    orig = {}
    for ref, part, fp, pn in wb[ORIG_SHEET].iter_rows(min_row=2, max_col=4, values_only=True):
        if ref:
            orig[str(ref).strip()] = (str(part or '').strip(), str(fp or '').strip(), str(pn or '').strip())
    picks = []
    for row in wb.worksheets[0].iter_rows(min_row=2):
        cells = row[:len(HEADER)]
        v = [str(c.value or '').strip() for c in cells] + [''] * len(HEADER)
        ref, pn = v[0], v[4]
        if not ref or ref not in orig or _is_summary(cells):
            continue
        part, fp, old = orig[ref]
        if pn and pn != old and pn not in (UNMATCHED, MECH_PN) and not NC_RE.search(part.upper()):
            picks.append((part, fp, pn, ref))
    return picks


def aggregate(details):
    """按料号重新分组：位号逗号连接、数量求和，按料号排序。"""
    groups = {}
    for ref, pn, name, spec, qty in details:
        g = groups.setdefault(pn, [[], 0, name, spec])
        g[0].append(ref)
        g[1] += int(qty)
    return [[','.join(g[0]), pn, g[2], g[3], str(g[1])]
            for pn, g in sorted(groups.items())]


def write_final(manual, aggregated, specials, leftovers, output_file):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(FINAL_HEADER)
    ws.freeze_panes = 'A2'

    thin = openpyxl.styles.Side(style='thin')
    border = openpyxl.styles.Border(left=thin, right=thin, top=thin, bottom=thin)
    yellow = openpyxl.styles.PatternFill('solid', start_color='FFFF00')

    seq = 0
    for item, is_left in ([(x, False) for x in manual + aggregated + specials]
                          + [(x, True) for x in leftovers]):
        ref, pn, name, spec, qty = item
        seq += 1
        row_num = ws.max_row + 1
        ws.append([seq, pn, name, spec, qty, ref])
        for col in range(1, len(FINAL_HEADER) + 1):
            cell = ws.cell(row=row_num, column=col)
            cell.border = border
            if is_left:
                cell.fill = yellow

    center = openpyxl.styles.Alignment(horizontal='center', vertical='center', wrap_text=True)
    left = openpyxl.styles.Alignment(horizontal='left', vertical='center', wrap_text=True)
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
        for cell in row:
            cell.alignment = center if (cell.row == 1 or cell.column in (1, 2, 3, 5)) else left
    for col in range(1, len(FINAL_HEADER) + 1):
        ws.cell(row=1, column=col).border = border

    widths = {1: 7, 2: 16, 3: 14, 4: 44, 5: 10, 6: 44}
    for col, w in widths.items():
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = w

    ws.row_dimensions[1].height = 22
    col_eff = {4: 36, 6: 36}
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        max_lines = 1
        for cell in row:
            if cell.column in col_eff and cell.value:
                w = sum(2 if ord(c) > 127 else 1 for c in str(cell.value))
                max_lines = max(max_lines, math.ceil(w / col_eff[cell.column]))
        ws.row_dimensions[row[0].row].height = max(30, max_lines * 15)

    wb.save(output_file)


def finalize(result_file, output_file=None, log=print, code_lookup=None, stopped=()):
    """入口：结果 Excel -> <原名>_生产版.xlsx，返回输出路径。
    code_lookup: 编码表 {料号: (名称, 规格)}；stopped: 停用料号集合（填了要警告）。"""
    if not output_file:
        base, ext = os.path.splitext(result_file)
        output_file = base + '_生产版' + (ext or '.xlsx')

    manual, details, specials, leftovers, n_filled = read_result(result_file, code_lookup)
    aggregated = aggregate(details)
    if n_filled:
        log(f"{n_filled} 行按料号从编码表补齐了名称/规格")
    filled = manual + details + specials
    bad_stop = [f"{x[0] or '(无位号)'}={x[1]}" for x in filled if x[1].upper() in stopped]
    if bad_stop:
        log(f"警告: {len(bad_stop)} 行填的是停用料，请换料: {'、'.join(bad_stop)}")
    stale = [f"{x[0] or '(无位号)'}={x[1]}" for x in filled
             if x[2] in ('', UNMATCHED) and x[1].upper() not in stopped]
    if stale:
        log(f"警告: {len(stale)} 行的料号在编码表里查不到、名称也没填，请核对料号: {'、'.join(stale)}")
    if leftovers:
        refs = '、'.join(x[0] or '(无位号)' for x in leftovers)
        log(f"警告: {len(leftovers)} 行未填料号，已保留黄色置于表尾: {refs}")

    write_final(manual, aggregated, specials, leftovers, output_file)
    log(f"生产版已生成: {output_file}（{len(manual)}行追加物料 + {len(aggregated)}行汇总 "
        f"+ {len(specials)}行特殊数量 + {len(leftovers)}行未填料号）")
    return output_file
