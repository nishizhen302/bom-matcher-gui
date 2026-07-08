# -*- coding: utf-8 -*-
"""加载层：BOM 与编码表 → 结构化列表。停用料在这里一次性剔除。"""
import re
import openpyxl

from .model import Entry, CAT_R, CAT_C, CAT_L, CAT_CONN, CAT_OTHER
from . import parsers as P

PN_RE = re.compile(r'^Y(?:06|12)\d{2}-\d+', re.I)


def load_bom(path):
    """返回 (rows, col_map)。rows为表头之后的原始行。表头行自动定位。"""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    for idx, row in enumerate(rows):
        header = [re.sub(r'\s+', ' ', str(c).strip().lower()) if c else '' for c in row]
        if 'reference' in header and 'part' in header and \
           ('pcb footprint' in header or 'footprint' in header):
            cols = {
                'ref': header.index('reference'),
                'part': header.index('part'),
                'fp': header.index('pcb footprint') if 'pcb footprint' in header else header.index('footprint'),
                'qty': header.index('quantity') if 'quantity' in header else (
                    header.index('qty') if 'qty' in header else None),
            }
            return rows[idx + 1:], cols
    raise ValueError("BOM表头未找到：需要包含 Reference、Part、PCB Footprint 或 Footprint")


def _category_of(pn, name):
    if pn.startswith('Y0601'):
        return CAT_R
    if pn.startswith('Y0602') or '电容' in name:
        return CAT_C
    if pn.startswith(('Y0603', 'Y0604', 'Y0606')) or '电感' in name or '磁珠' in name:
        return CAT_L
    if pn.startswith('Y0614') or '连接器' in name:
        return CAT_CONN
    return CAT_OTHER


def load_code(path):
    """返回 (entries, discontinued, sheet_names)。entries 已剔除停用料并完成解析。"""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    entries, discontinued, sheet_names = [], [], []
    for sheet in wb.sheetnames:
        found = False
        for row in wb[sheet].iter_rows(values_only=True):
            cells = list(row)
            pn_idx = next((i for i, c in enumerate(cells)
                           if c and PN_RE.match(str(c).strip())), None)
            if pn_idx is None:
                continue
            found = True
            pn = str(cells[pn_idx]).strip()
            name = str(cells[pn_idx + 1]).strip() if pn_idx + 1 < len(cells) and cells[pn_idx + 1] else ''
            spec = str(cells[pn_idx + 2]).strip() if pn_idx + 2 < len(cells) and cells[pn_idx + 2] else ''
            e = _build_entry(pn, name, spec)
            if e.discontinued:
                discontinued.append(e)
            else:
                entries.append(e)
        if found:
            sheet_names.append(sheet)
    wb.close()
    if not entries:
        raise ValueError("编码表未找到 Y06/Y12 料号行")
    entries.sort(key=lambda e: e.pn)  # 匹配结果与表内行序无关
    return entries, discontinued, sheet_names


def _build_entry(pn, name, spec):
    cat = _category_of(pn, name)
    e = Entry(pn=pn, name=name, spec=spec, category=cat,
              discontinued=('停用' in spec or '停用' in name),
              tokens=tuple(P.tokens_of(spec)), norm=P.norm(spec))
    if cat == CAT_R:
        e.r_ohms, e.package, e.tol, e.is_array = P.parse_spec_r(spec)
        e.is_array = e.is_array or ('排阻' in name)
    elif cat == CAT_C:
        e.c_pf, e.volt, e.package = P.parse_spec_c(spec)
        e.dims = P.parse_dims(spec)
    elif cat == CAT_L:
        e.l_uh = P.parse_spec_l(spec)
        e.package, _ = P.parse_package(spec, spec_mode=True)
        e.dims = P.parse_dims(spec)
    else:
        e.package, _ = P.parse_package(spec, spec_mode=True)
        e.volt = P.parse_volt(spec)
    return e
