# -*- coding: utf-8 -*-
"""规则层：替代/别名规则、联动规则、磁珠封装映射。

内置规则从旧版 SPECIAL_RULES 迁移而来。skill 目录下若存在 rules.xlsx
（工作表：替代 / 禁用，列：Part包含 | 封装包含 | 料号 | 备注），会追加在内置规则之前，
硬件工程师可直接编辑 Excel 调整规则，无需改代码。
"""
import os
import openpyxl

# (part包含, fp包含或None, [料号, 停用时的后备料号...])
ALIAS_RULES = [
    ('90R', 'CM_FILTER', ['Y0604-00030']),
    ('DVI', None, ['Y0614-00013']),
    ('DISPLAY_PORT-CON', None, ['Y0614-00018']),
    ('DP', 'DP_MOLEX_47272A', ['Y0614-00018']),
    ('GREEN', 'LED0603', ['Y0608-00007']),
    ('24C02', None, ['Y0612-00014']),
    ('24LC02', None, ['Y0612-00014']),
    ('X24C16', None, ['Y0612-00018']),   # 用户确认：用256K替代16K
    ('RTD2785', None, ['Y0612-00106']),
    ('RTD2795XX-CG', None, ['Y0612-00108']),
    ('RTD2797', None, ['Y0612-00118']),
    ('BGA_32X32', 'RTD2797', ['Y0612-00118']),  # Part与封装列颠倒的情况
    ('RTD2796', None, ['Y0612-00027']),
    ('RT9026', None, ['Y0612-00058']),
    ('HEATSINK', None, ['Y1209-00005']),
    ('RE-FUSE/49FKA-002', 'F1812', ['Y0607-00008']),
    ('47UF/35V', 'AE660', ['Y0602-00066']),
    ('100UF/25V', 'AE660', ['Y0602-00007']),
    ('100UF/50V', 'AE840', ['Y0602-00033']),
    ('100NF/0402', 'C0402', ['Y0602-00068']),
    ('100NF', 'C0402', ['Y0602-00068']),
    ('2N3904', None, ['Y0609-00001']),
    ('MMBT3904', None, ['Y0609-00001']),
    ('3904', None, ['Y0609-00001']),
    ('C4201WR', None, ['Y0614-00092']),
    ('UDZ3.6B', None, ['Y0608-00008', 'Y0608-00019']),  # 00008停用时自动切00019
]

# 内置禁用：编码表里重复登记的料，只留一个
BUILTIN_BANNED = {
    'Y0614-00027',   # 与 Y0614-00035 同为 B2B-PH-K-S，2026-10-18 用户确认停用 00027
}

# 磁珠：识别为磁珠后按封装直接映射（阻抗值不参与，与旧版策略一致）
BEAD_BY_PACKAGE = {
    '0603': 'Y0606-00006',
    '0805': 'Y0606-00003',
    '1206': 'Y0606-00012',
}

# 联动：匹配到 key 料号时自动追加 (料号, 每个的用量g, 说明)
COMPANIONS = {
    'Y1209-00005': ('Y1214-00006', 0.2, '散热硅脂(0.2g/散热器)'),
}

# SFLASH 条件规则：BOM 含 RTD2556 时用前者，否则后者
SFLASH_WITH_2556 = 'Y0612-00041'
SFLASH_DEFAULT = 'Y0612-00056'


def load_user_rules(skill_dir):
    """读取 rules.xlsx（若存在），返回 (extra_alias, banned_pns)。"""
    path = os.path.join(skill_dir, 'rules.xlsx')
    extra, banned = [], set()
    if not os.path.exists(path):
        return extra, banned
    wb = openpyxl.load_workbook(path, data_only=True)
    if '替代' in wb.sheetnames:
        for row in wb['替代'].iter_rows(min_row=2, values_only=True):
            if not row or not row[0] or not (len(row) > 2 and row[2]):
                continue
            part_kw = str(row[0]).strip().upper()
            fp_kw = str(row[1]).strip().upper() if row[1] else None
            pns = [p.strip() for p in str(row[2]).split(';') if p.strip()]
            extra.append((part_kw, fp_kw, pns))
    if '禁用' in wb.sheetnames:
        for row in wb['禁用'].iter_rows(min_row=2, values_only=True):
            if row and row[0]:
                banned.add(str(row[0]).strip())
    wb.close()
    return extra, banned
